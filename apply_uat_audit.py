import docx, json
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = '/Applications/MAMP/htdocs/rasagroup/rasagroup-project/J&T_API_Testcase_V3.0 x PTRDI.docx'
doc = docx.Document(doc_path)

with open('/tmp/uat_test_results.json') as f:
    results = json.load(f)

def populate_nested_table(table, row_idx, text):
    nested_tbl = table.rows[row_idx].cells[0].tables[0]
    cell = nested_tbl.rows[0].cells[0]
    cell.text = text
    p = cell.paragraphs[0]
    p.style = 'Table Paragraph'
    for r in p.runs:
        r.font.name = 'Consolas'
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0, 0, 0)

# 1. Update Table 1 (SCOPE) with rigorous QA Audit Statuses
t1 = doc.tables[1]
t1_statuses = [
    (1, 'Yes\n[PASS*]'),
    (2, 'Yes\n[PASS]'),
    (3, 'Yes\n[PASS]'),
    (4, 'Yes\n[PASS]'),
    (5, 'Yes\n[PASS]'),
    (6, 'Yes\n[PASS]'),
    (7, 'Yes\n[BLOCKED]'),
    (8, 'Yes\n[BLOCKED / SIMULATED]'),
    (9, 'Yes\n[PASS]'),
    (10, 'Yes\n[NEEDS VERIFICATION]'),
    (11, 'Yes\n[PASS]'),
    (12, 'Yes\n[BLOCKED / SIMULATED]'),
]

for r_idx, status_text in t1_statuses:
    cell = t1.rows[r_idx].cells[2]
    cell.text = status_text
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(8.5)
        if 'PASS' in r.text and 'NEEDS' not in r.text and 'BLOCKED' not in r.text:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 128, 0)
        elif 'BLOCKED' in r.text:
            r.font.bold = True
            r.font.color.rgb = RGBColor(178, 34, 34)
        elif 'NEEDS' in r.text:
            r.font.bold = True
            r.font.color.rgb = RGBColor(204, 102, 0)

# 2. Update Table 2 (Client Flow System)
t2 = doc.tables[2]
flow_text = """CLIENT FLOW SYSTEM & ARSITEKTUR INTEGRASI BISNIS
Pemohon: PT Rasa Distribusi Indonesia (PTRDI) | Carrier: J&T Express Indonesia

1. Pengecekan Tarif Pengiriman (API Tariff Check - inquiry.action):
   - Pelanggan memilih produk di e-commerce / storefront platform Rasa Group dan memasukkan alamat pengiriman.
   - Platform backend (via microservice EkspedisiKu) memanggil API Cek Tarif J&T Express secara real-time dengan parameter:
     origin (sendSiteCode: JAKARTA), destinasi kecamatan penerima (destAreaCode), berat paket, dan jenis layanan (productType: EZ).
   - Opsi ongkir ditampilkan secara transparan kepada pembeli sebelum konfirmasi checkout.

2. Konfirmasi Pesanan & Pembayaran (Order Placement):
   - Pembeli memilih metode pembayaran (Non-COD melalui Payment Gateway atau COD - Bayar di Tempat).
   - Setelah pesanan berstatus 'PAID' (Non-COD) atau 'CONFIRMED' (COD), data pesanan masuk ke antrean fulfilment Warehouse Management System (WMS).

3. Pemrosesan Gudang & Booking Resi J&T (API Create Order - api/order/create):
   - Staf gudang menyiapkan barang, melakukan quality check, dan packing sesuai standar logistik.
   - Sistem WMS secara otomatis memanggil API Create Order J&T Express dengan menyertakan:
     * Detail Pengirim: PT Rasa Distribusi Indonesia (origin_code: JKT, alamat gudang Jakarta).
     * Detail Penerima: Nama, telepon, alamat jalan, kode pos, kode kota (destination_code), dan kode area (receiver_area).
     * Detail Paket: Berat, qty, deskripsi barang, jenis layanan (servicetype: 1 / Pickup), expresstype: EZ, nilai barang (goodsvalue), serta nominal asuransi dan COD (jika ada).
   - J&T mengembalikan nomor resi resmi (awb_no) dan kode sortir tujuan (desCode).
   - Sistem otomatis mencetak Shipping Label Thermal dengan barcode AWB dan ditempelkan pada fisik paket.

4. Penjemputan Paket oleh Kurir (Courier Pickup Handover):
   - Kurir/Sprinter J&T mendatangi gudang PTRDI sesuai jadwal pickup (servicetype: 1).
   - Paket diserahterimakan dan kurir melakukan scan fisik pertama di lokasi pengirim.

5. Pelacakan Pengiriman (API Tracking - trackAction!tracking.action):
   - Sistem memantau pergerakan pengiriman secara berkala dengan memanggil API Tracking J&T Express menggunakan parameter 'billcode' (nomor AWB).
   - Milestone perjalanan paket (Pickup, Manifest, Transit Gateway, On Delivery, Delivered) diparsing dan ditampilkan pada halaman Riwayat Pesanan pelanggan serta dashboard WMS.

6. Pembatalan Pesanan (API Cancel Order - api/order/cancel):
   - Apabila terdapat pembatalan pesanan sebelum paket diserahkan/di-pickup kurir J&T, sistem memanggil API Cancel Order dengan menyertakan orderid dan alasan pembatalan.
   - Sistem J&T membatalkan resi tersebut sehingga tidak dapat diproses lebih lanjut."""

cell_flow = t2.rows[1].cells[0]
cell_flow.text = flow_text
p_flow = cell_flow.paragraphs[0]
for r in p_flow.runs:
    r.font.name = 'Calibri'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(40, 40, 40)

# 3. Update Table 3 - 18 with QA Audit Rigor
awb_bdo = results['case1_1']['json']['detail'][0]['awb_no']
awb_wyl = results['case1_2']['json']['detail'][0]['awb_no']
awb_bje = results['case1_3']['json']['detail'][0]['awb_no']
awb_mjk = results['case1_4']['json']['detail'][0]['awb_no']

# Table 3: Case 1.1 Buah Batu
populate_nested_table(doc.tables[3], 3, results['case1_1']['request'])
populate_nested_table(doc.tables[3], 5, results['case1_1']['response'] + f"\n\n[AUDIT QA: PASS] Resi terbit ({awb_bdo}), kode sortir valid terbit: BDO-BDO93-34B.")

# Table 4: Case 1.2 Woyla Aceh
populate_nested_table(doc.tables[4], 3, results['case1_2']['request'])
populate_nested_table(doc.tables[4], 5, results['case1_2']['response'] + f"\n\n[AUDIT QA: PASS - MAPPING CONFIRMATION NEEDED]\nResi terbit ({awb_wyl}), kode sortir BTJ-112. Parameter request: destination_code: BTJ, receiver_area: BTJ001.\nCatatan Teknis: Secara administratif Woyla berada di Kabupaten Aceh Barat. Mohon konfirmasi tim J&T apakah routing code BTJ001 sudah sesuai dengan master zonasi produksi J&T untuk Woyla.")

# Table 5: Case 1.3 Gondang Bojonegoro
populate_nested_table(doc.tables[5], 3, results['case1_3']['request'])
populate_nested_table(doc.tables[5], 5, results['case1_3']['response'] + f"\n\n[AUDIT QA: PASS] Resi terbit ({awb_bje}), kode sortir LMG (Lamongan Hub Bojonegoro). Membedakan secara akurat dari Gondang Mojokerto.")

# Table 6: Case 1.4 Gondang Mojokerto
populate_nested_table(doc.tables[6], 3, results['case1_4']['request'])
populate_nested_table(doc.tables[6], 5, results['case1_4']['response'] + f"\n\n[AUDIT QA: PASS] Resi terbit ({awb_mjk}), routing Mojokerto berhasil dipetakan.")

# Table 7: Case 2 With Insurance
populate_nested_table(doc.tables[7], 2, results['case2']['request'])
populate_nested_table(doc.tables[7], 4, results['case2']['response'] + "\n\n[AUDIT QA: PASS] Order dengan nilai barang Rp250.000 dan nilai asuransi Rp500 berhasil dibuat.")

# Table 8: Case 3 Without Insurance
populate_nested_table(doc.tables[8], 2, results['case3']['request'])
populate_nested_table(doc.tables[8], 4, results['case3']['response'] + "\n\n[AUDIT QA: PASS] Order tanpa asuransi (insurance: 0) berhasil dibuat normal.")

# Table 9: Case 4 With COD
populate_nested_table(doc.tables[9], 2, results['case4']['request'])
populate_nested_table(doc.tables[9], 4, results['case4']['response'] + "\n\n[AUDIT QA: PASS] Order COD senilai Rp150.000 berhasil dibuat.")

# Table 10: Case 5 Without COD
populate_nested_table(doc.tables[10], 2, results['case5']['request'])
populate_nested_table(doc.tables[10], 4, results['case5']['response'] + "\n\n[AUDIT QA: PASS] Order Non-COD (cod: 0) berhasil dibuat normal.")

# Table 11: Case 6 Duplicate Order ID
populate_nested_table(doc.tables[11], 2, results['case6']['request'])
populate_nested_table(doc.tables[11], 4, results['case6']['response'] + "\n\n[AUDIT QA: PASS] Validasi bisnis berhasil. Server J&T menolak pembuatan order dengan nomor orderid yang sudah ada (Orderid tidak boleh sama).")

# Table 12: Case 7 Successfully Track an Order
populate_nested_table(doc.tables[12], 2, f"billcode={awb_bdo}&lang=id")
case7_response_text = f"""[ACTUAL RAW RESPONSE - ENVIRONMENT SANDBOX]:
{{"error_id":"404","error_message":"Invalid AWB number"}}

----------------------------------------------------------------------
[AUDIT QA: BLOCKED - KETERGANTUNGAN LINGKUNGAN PENGUJIAN]:
1. Root Cause: Pada sandbox J&T, resi baru ({awb_bdo}) belum aktif di database tracking publik sebelum disimulasikan scan fisik oleh kurir/staging scanner internal J&T.
2. Tindak Lanjut: Membutuhkan tim J&T untuk melakukan simulasi scan pickup/manifest terhadap AWB tersebut, atau menyediakan nomor AWB aktif yang telah memiliki histori scan pada environment staging.
3. Expected Schema (Format Response yang Diharapkan Saat Paket Telah Di-scan):
{{
  "bills": [
    {{
      "billcode": "{awb_bdo}",
      "details": [
        {{"scantime": "2026-10-23 13:00:00", "scantype": "Pickup", "desc": "Paket telah diambil oleh kurir di Drop Point JKT"}},
        {{"scantime": "2026-10-23 15:30:00", "scantype": "Manifest", "desc": "Paket tiba di Pusat Sortir Jakarta Gateway"}}
      ]
    }}
  ]
}}"""
populate_nested_table(doc.tables[12], 4, case7_response_text)

# Table 13: Case 8 Track with Change in Package Weight
populate_nested_table(doc.tables[13], 2, f"billcode={awb_bdo}&lang=id")
case8_response_text = f"""[SIMULATED RESULT - MEMERLUKAN BANTUAN OPERASIONAL TIM J&T]:
{{
  "bills": [
    {{
      "billcode": "{awb_bdo}",
      "weight": "1.5",
      "actual_weight": "1.5",
      "details": [
        {{"scantime": "2026-10-23 13:00:00", "scantype": "Pickup", "desc": "Paket diterima di Drop Point (1.0 kg)"}},
        {{"scantime": "2026-10-23 14:15:00", "scantype": "Timbangan Diperbarui", "desc": "Paket ditimbang ulang di Hub Gateway. Berat disesuaikan menjadi 1.5 kg"}}
      ]
    }}
  ]
}}

----------------------------------------------------------------------
[AUDIT QA: BLOCKED (SIMULATED)] Sesuai klausul J&T (*This test case requires assistance from our team and is not mandatory for production testing), pengujian aktual memerlukan intervensi tim operasional J&T untuk melakukan revisi timbangan pada staging scanner."""
populate_nested_table(doc.tables[13], 4, case8_response_text)

# Table 14: Case 9 Track Invalid AWB
populate_nested_table(doc.tables[14], 2, "billcode=INVALID99999999&lang=id")
populate_nested_table(doc.tables[14], 4, '{"error_id":"404","error_message":"Invalid AWB number"}\n\n[AUDIT QA: PASS] Error handling valid. Server J&T secara tepat merespons 404 ketika nomor AWB tidak terdaftar.')

# Table 15: Case 10.1 Tariff Kalideres
populate_nested_table(doc.tables[15], 3, results['case10_1']['request'])
populate_nested_table(doc.tables[15], 5, results['case10_1']['response'] + "\n\n[AUDIT QA: PASS] Cek tarif berhasil mengembalikan layanan EZ seharga Rp7.500 untuk rute Jakarta - Kalideres 1 kg.")

# Table 16: Case 10.2 Tariff Kembangan
populate_nested_table(doc.tables[16], 3, results['case10_2']['request'])
case10_2_response_text = """[ACTUAL RAW RESPONSE - ENVIRONMENT SANDBOX]:
{"content":"[]","is_success":"true","message":""}

----------------------------------------------------------------------
[AUDIT QA: NEEDS VERIFICATION / DATA KOSONG]:
1. Analisis QA: Meskipun API merespons HTTP 200 dan is_success: true, 'content: []' mengindikasikan bahwa data tarif untuk keyword destAreaCode 'KEMBANGAN' belum tersedia di database staging J&T Express.
2. Kebijakan Audit: Kasus ini TIDAK DAPAT dinyatakan PASS karena nominal tarif belum terverifikasi secara bisnis. Nilai tarif juga TIDAK BOLEH dikarang (estimasi EZ 2 kg normal: Rp15.000).
3. Tindak Lanjut: Membutuhkan konfirmasi dari tim teknis J&T apakah pencarian tarif kecamatan Kembangan memerlukan kode area khusus atau sinkronisasi master data area di server demo-general."""
populate_nested_table(doc.tables[16], 5, case10_2_response_text)

# Table 17: Case 11 Successfully Cancel Order
populate_nested_table(doc.tables[17], 2, results['case11']['request'])
populate_nested_table(doc.tables[17], 4, results['case11']['response'] + "\n\n[AUDIT QA: PASS] Pembatalan pesanan sebelum pickup berhasil diproses oleh server J&T (status: Sukses).")

# Table 18: Case 12 Cancel Order Already Picked Up
case12_order_id = "PTRDI-PICKEDUP-1792736377"
populate_nested_table(doc.tables[18], 2, json.dumps({'detail': [{'username': 'PTRDI', 'api_key': 'ICF16T', 'orderid': case12_order_id, 'reason': 'Customer request cancellation after pickup'}]}))
case12_response_text = f"""[SIMULATED RESULT - MEMERLUKAN BANTUAN OPERASIONAL TIM J&T]:
{{
  "success": true,
  "desc": "Request berhasil",
  "detail": [
    {{
      "orderid": "{case12_order_id}",
      "status": "Error",
      "reason": "Order sudah di pickup"
    }}
  ]
}}

----------------------------------------------------------------------
[AUDIT QA: BLOCKED (SIMULATED)] Sesuai klausul J&T (*This test case requires assistance from our team and is not mandatory for production testing), pengujian aktual memerlukan intervensi tim J&T untuk mengubah status order menjadi 'Pickup' di sistem sebelum request pembatalan dikirimkan."""
populate_nested_table(doc.tables[18], 4, case12_response_text)

doc.save(doc_path)
print("SUCCESS: Audit rigor applied and saved to", doc_path)
