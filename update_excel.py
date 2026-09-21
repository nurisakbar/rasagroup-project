import openpyxl
import json

file_path = '/Users/nurisakbar/Downloads/FASPAY QRIS - Skenario Functional Test_V.3.2.xlsx'
wb = openpyxl.load_workbook(file_path)
sheet = wb.active

req_text = """URL: https://debit-sandbox.faspay.co.id/v1.0/qr/qr-mpm-generate

HEADER:
{
  "X-TIMESTAMP": "2026-09-09T23:32:01+07:00",
  "X-SIGNATURE": "lAtSoraQPatUUYphBuc6JIXK0Fh5wXUxSmuU2o29zZnMt4WzEMrIb5iL7uOz8ulhujimgpsl1W0CJ49hOzQRVyjmD9nQxoEh5izAJ9KZEp2rd8l85gHL/eJo8vxbjiMYALfxWkvM6rUEk0hAGhX/ic5FaHk5ZGlTe2Ele+WkxAHOMPW66QSi1pvDeFNuUMt8diLpR7eLydTYWS1DPgq9u7G5txfW2tz7/oOpV7oNbIE3op6N3MPWYcO2fkipe9kNHtDd+jo068j8j+g4KJMDUFM2CPPkKJ5FqWM7MFAHS2vQzewAEK/Jt3MWqGDmVgNzRPIwIg8n56Cx3WJp8/8DdA==",
  "X-PARTNER-ID": "37116",
  "X-EXTERNAL-ID": "202609091632016515",
  "CHANNEL-ID": "77001",
  "Content-Type": "application/json"
}

BODY:
{
  "partnerReferenceNo": "WS260909018",
  "amount": {
    "value": "1000.00",
    "currency": "IDR"
  },
  "merchantId": "37116",
  "validityPeriod": "2026-09-10T00:32:01+07:00",
  "additionalInfo": {
    "billDate": "2026-09-09T23:32:01+07:00",
    "billDescription": "Payment #WS260909018",
    "channelCode": "711",
    "phoneNo": "089699935552"
  }
}"""

res_text = """BODY:
{
  "responseCode": "2004700",
  "responseMessage": "Request has been processed successfully",
  "referenceNo": "3711671152289727",
  "partnerReferenceNo": "WS260909018",
  "qrContent": "00020101021226560016ID.CO.SHOPEE.WWW011893600918000000026602032660303UBE52041234530336054071000.005802ID5912Faspay Store6015KOTA JAKARTA SE6105123456222051811856332618570381463049E90",
  "qrUrl": "https://debit-sandbox.faspay.co.id/__assets/qr/37116-3711671152289727.png",
  "additionalInfo": {
    "qrImageUrl": "https://debit-sandbox.faspay.co.id/pws/100003/0830000010100000/ee3d6870ae94b0ad93ecf7817c37343b77d9d832?trx_id=3711671152289727&merchant_id=37116&bill_no=WS260909018",
    "merchantId": "37116",
    "amount": {
      "value": "1000.00",
      "currency": "IDR"
    },
    "phoneNo": "089699935552"
  }
}"""

# Find row 18.6
for row in range(1, 100):
    cell_val = str(sheet.cell(row=row, column=1).value).strip()
    if cell_val == '18.6':
        sheet.cell(row=row, column=5).value = req_text  # Request
        sheet.cell(row=row, column=6).value = res_text  # Response
        sheet.cell(row=row, column=7).value = 'PASS'    # Result
        break

new_file_path = '/Users/nurisakbar/Downloads/FASPAY QRIS - Skenario Functional Test_V.3.2_Updated_Valid.xlsx'
wb.save(new_file_path)
print(f"Saved to {new_file_path}")
